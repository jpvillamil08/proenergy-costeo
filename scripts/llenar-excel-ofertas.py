#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Llena el Excel de seguimiento de ofertas con los datos de la aplicacion.

Reglas acordadas con PROENERGY:
  - Lo que YA esta escrito en el Excel manda. Solo se rellenan las celdas
    vacias; nunca se pisa un dato existente, aunque la app diga otra cosa
    (hay filas donde el cliente o el precio difieren de lo importado de Siigo,
    y se asume que la version revisada a mano es la buena).
  - Se incluyen todas las cotizaciones de 2026 hasta hoy, no solo las que ya
    estaban listadas.
  - Las tres columnas de costo se reparten por tipo de linea:
      COSTO PROVEEDORES ....... materiales
      VIATICOS + GASOLINA ..... carrocanasta, grua, transporte
      MO PROENERGY ............ horas de tecnicos y linieros
  - CIUDAD se deduce del texto de la cotizacion cuando se puede; el resto de
    campos que la app no tiene (direccion, tiempo de ejecucion, validez de la
    oferta) quedan vacios.

Uso:
    $env:ADMIN_USERNAME = "admin"
    $env:ADMIN_PASSWORD = "tu-clave"
    python scripts/llenar-excel-ofertas.py <URL> "<ruta del excel>"
"""

import importlib.util
import os
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

_d = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location('inventario', _d / 'listar-elementos-cotizados.py')
inv = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(inv)

HOY = date.today().isoformat()
ANIO = HOY[:4]


def norm(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s.upper()).strip()


# ---- clasificacion de cada linea de costo -------------------------------
# Se reparte el costo en las tres columnas del Excel segun lo que sea la linea.
VIATICOS = ('CARROCANASTA', 'CARRO CANASTA', 'GRUA', 'TRANSPORTE', 'ALQUILER', 'YALE')
MANO_OBRA = ('MANO DE OBRA', 'APERTURA', 'CIERRE', 'RETIRO E INSTALACION',
             'RETIRO Y/O INSTALACION', 'INSTALACION DE MEDIDOR', 'PREVISITA',
             'VISITA TECNICA', 'MANTENIMIENTO', 'LIMPIEZA', 'GESTION', 'MONTAJE',
             'CAMBIO DE DPS', 'DESCONEXION', 'RECONEXION', 'CALIBRACION',
             'MANIOBRA', 'CONEXIONADO', 'SUPERVISION', 'TERMOGRAFIA', 'PODA',
             'ATENCION A EMERGENCIA', 'INTERVENCION', 'ADECUACION', 'OBRAS ELECTRICAS')


def clasificar_linea(descripcion):
    d = norm(descripcion).replace('[REVISAR]', '').strip()
    if any(d.startswith(p) or f' {p}' in d for p in VIATICOS):
        return 'viaticos'
    if any(d.startswith(p) for p in MANO_OBRA):
        return 'mano_obra'
    return 'proveedores'


# ---- ciudad, deducida del texto -----------------------------------------
CIUDADES = [
    'BARRANQUILLA', 'CARTAGENA', 'SANTA MARTA', 'SINCELEJO', 'MONTERIA',
    'VALLEDUPAR', 'RIOHACHA', 'SOLEDAD', 'MALAMBO', 'GALAPA', 'PUERTO COLOMBIA',
    'SABANALARGA', 'LURUACO', 'POZOS COLORADOS', 'PIVIJAY', 'BOGOTA', 'MEDELLIN',
    'CIENAGA', 'FUNDACION', 'PLATO', 'MAGANGUE', 'COROZAL', 'TOLU', 'BARANOA',
    'PALMAR DE VARELA', 'PONEDERA', 'REPELON', 'CANDELARIA', 'CAMPO DE LA CRUZ',
    'TURBACO', 'ARJONA', 'MAMONAL', 'RODADERO', 'GAIRA',
]


def deducir_ciudad(*textos):
    t = norm(' '.join(str(x or '') for x in textos))
    for ciudad in CIUDADES:
        if ciudad in t:
            return ciudad
    return None


# ---- tipo de actividad, deducido ----------------------------------------
ACTIVIDADES = [
    ('PREVISITAS', ('PREVISITA',)),
    ('MANTENIMIENTO', ('MANTENIMIENTO',)),
    ('CARRO CANASTA', ('CARROCANASTA', 'CARRO CANASTA')),
    ('CAMBIO DE EQUIPO DE MEDIDA', ('GRUPO DE MEDIDA', 'EQUIPO DE MEDIDA', 'MEDIDA DIRECTA',
                                    'MEDIDA INDIRECTA', 'MEDIDA SEMIDIRECTA', 'MEDIDOR')),
    ('CALIBRACION', ('CALIBRACION',)),
    ('PRUEBAS ELECTRICAS', ('PRUEBA', 'TERMOGRAFIA', 'ANALIZADOR DE RED')),
    ('SUMINISTRO DE MATERIALES', ('SUMINISTRO', 'MATERIAL')),
    ('OBRA ELECTRICA', ('OBRA', 'ADECUACION', 'INSTALACION', 'MONTAJE', 'CONEXION')),
    ('GESTION ANTE OR', ('GESTION', 'DESCARGO')),
]


def deducir_actividad(*textos):
    t = norm(' '.join(str(x or '') for x in textos))
    for etiqueta, claves in ACTIVIDADES:
        if any(k in t for k in claves):
            return etiqueta
    return None


COL = {}  # se llena leyendo los encabezados del Excel


def main():
    if len(sys.argv) < 3:
        print('Uso: python scripts/llenar-excel-ofertas.py <URL> "<ruta del excel>"', file=sys.stderr)
        sys.exit(1)
    ruta = Path(sys.argv[2])
    if not ruta.exists():
        print(f'No existe el archivo: {ruta}', file=sys.stderr)
        sys.exit(1)
    cfg = inv.parsear_argumentos([sys.argv[1]])
    usuario, clave = os.environ.get('ADMIN_USERNAME'), os.environ.get('ADMIN_PASSWORD')
    if not usuario or not clave:
        print('Faltan ADMIN_USERNAME / ADMIN_PASSWORD.', file=sys.stderr)
        sys.exit(1)

    api = inv.ClienteApi(cfg['base_url'])
    api.login(usuario, clave)
    print('Leyendo la aplicacion...')

    lista = api.get('/api/cotizaciones') or []
    facturas = api.get('/api/facturas?desde=2000-01-01&hasta=2030-12-31') or []
    fac_por_cot = {}
    for f in facturas:
        if f.get('cotizacion_id'):
            fac_por_cot.setdefault(f['cotizacion_id'], []).append(f)

    del_anio = [c for c in lista
                if str(c.get('fecha_cotizacion') or '').startswith(ANIO)
                and str(c.get('fecha_cotizacion'))[:10] <= HOY]
    print(f'Cotizaciones de {ANIO} hasta {HOY}: {len(del_anio)}')

    datos = {}
    for i, c in enumerate(del_anio, 1):
        try:
            det = api.get(f"/api/cotizaciones/{c['id']}")
        except RuntimeError:
            continue
        cp = det['calculo']['costeoPresupuestado']
        reparto = {'proveedores': 0.0, 'viaticos': 0.0, 'mano_obra': 0.0}
        for m in det.get('materiales') or []:
            valor = float(m.get('cantidad_presupuestada') or 0) * float(m.get('costo_unitario') or 0)
            reparto[clasificar_linea(m.get('descripcion'))] += valor
        # La mano de obra cargada como tal (hoy no hay, pero si la agregan cuenta)
        for mo in det.get('manoObra') or []:
            factor = float(mo.get('factor_prestacional') or 1) if mo.get('tipo') == 'Interno' else 1
            reparto['mano_obra'] += float(mo.get('horas_presupuestadas') or 0) * float(mo.get('tarifa_hora') or 0) * factor

        facs = fac_por_cot.get(c['id'], [])
        orden = next((f.get('orden') for f in facs if f.get('orden')), None)
        datos[norm(c['numero'])] = {
            'numero': c['numero'],
            'fecha': str(c.get('fecha_cotizacion') or '')[:10],
            'cliente': c.get('cliente'),
            'estado': c.get('estado'),
            'proyecto': (c.get('descripcion') or '').strip(),
            'ciudad': deducir_ciudad(c.get('descripcion'), c.get('cliente')),
            'actividad': deducir_actividad(c.get('descripcion')),
            'proveedores': round(reparto['proveedores'], 2),
            'viaticos': round(reparto['viaticos'], 2),
            'mano_obra': round(reparto['mano_obra'], 2),
            'imprevistos': round(float(cp.get('imprevistos') or 0), 2),
            'utilidad': round(float(cp.get('utilidad') or 0), 2),
            'margen': round(float(cp.get('margenPct') or 0), 4),
            'precio': float(c.get('precio_venta') or 0),
            'anticipo': float(c.get('pct_anticipo') or 0),
            'plazo': int(c.get('dias_credito_otorgados') or 0),
            'facturas': ', '.join(f['numero'] for f in facs) if facs else None,
            'orden': orden,
        }
        if i % 25 == 0 or i == len(del_anio):
            print(f'  {i}/{len(del_anio)}...', end='\r', flush=True)
    print()
    escribir(ruta, datos)


def escribir(ruta, datos):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    wb = load_workbook(ruta)
    ws = wb[wb.sheetnames[0]]
    FILA_ENC = 2

    # Mapa encabezado -> columna, leyendo lo que el Excel ya tiene
    enc = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=FILA_ENC, column=c).value
        if v:
            enc[norm(v)] = c

    def col(*nombres):
        for n in nombres:
            for k, v in enc.items():
                if k.startswith(norm(n)):
                    return v
        return None

    C = {
        'numero': col('OFERTA No'), 'fecha': col('FECHA DE ENVIO'), 'cliente': col('CLIENTE'),
        'estado': col('ESTADO'), 'proyecto': col('NOMBRE DEL PROYECTO'), 'ciudad': col('CIUDAD'),
        'actividad': col('TIPO DE ACTIVIDAD'), 'proveedores': col('COSTO PROVEEDORES'),
        'viaticos': col('COSTO VIATICOS'), 'mano_obra': col('COSTO MO'),
        'margen': col('UTILIDAD %'), 'utilidad': col('UTILIDAD $'), 'precio': col('PV (SIN IVA)', 'PV'),
        'anticipo': col('ANTICIPO'), 'imprevistos': col('IMPREVISTOS'),
        'plazo': col('PLAZO DE PAGO'), 'facturas': col('NUMERO DE FACTURA'),
        'orden': col('OBSERVACIONES'),
    }
    faltan = [k for k, v in C.items() if v is None]
    if faltan:
        print(f'  aviso: no se encontro columna para {faltan}')

    # Filas que ya existen, por numero de oferta
    fila_de = {}
    ultima = FILA_ENC
    for r in range(FILA_ENC + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=C['numero']).value
        if v:
            fila_de[norm(v)] = r
            ultima = r

    NUEVA = PatternFill('solid', fgColor='EAF3FF')   # filas que agrega el script
    RELLENO = Font(color='1F4E79')                    # celdas que rellena en filas existentes

    rellenadas = agregadas = respetadas = 0
    for clave, d in sorted(datos.items(), key=lambda x: x[1]['fecha'] or ''):
        nueva = clave not in fila_de
        if nueva:
            ultima += 1
            r = ultima
            ws.cell(row=r, column=C['numero']).value = d['numero']
        else:
            r = fila_de[clave]

        for campo, columna in C.items():
            if columna is None or campo == 'numero':
                continue
            valor = d.get(campo)
            if valor in (None, '', 0) and campo not in ('anticipo', 'plazo'):
                continue
            celda = ws.cell(row=r, column=columna)
            if celda.value not in (None, ''):
                respetadas += 1
                continue          # lo que ya escribieron a mano manda
            celda.value = valor
            if not nueva:
                celda.font = RELLENO
                rellenadas += 1
            if campo in ('proveedores', 'viaticos', 'mano_obra', 'utilidad', 'precio', 'imprevistos'):
                celda.number_format = '#,##0'
            elif campo == 'margen':
                celda.number_format = '0.0%'
            elif campo == 'fecha':
                celda.number_format = 'dd/mm/yyyy'
        if nueva:
            agregadas += 1
            for cc in range(3, max(C.values()) + 1):
                if ws.cell(row=r, column=cc).fill.fgColor.rgb in (None, '00000000'):
                    ws.cell(row=r, column=cc).fill = NUEVA

    salida = ruta.with_name(ruta.stem + ' - LLENO.xlsx')
    wb.save(salida)
    print(f'\nGuardado: {salida}')
    print(f'  filas nuevas agregadas      : {agregadas}')
    print(f'  celdas vacias rellenadas    : {rellenadas}')
    print(f'  celdas respetadas (ya tenian): {respetadas}')


if __name__ == '__main__':
    try:
        main()
    except RuntimeError as e:
        print(f'\nERROR: {e}', file=sys.stderr)
        sys.exit(1)
