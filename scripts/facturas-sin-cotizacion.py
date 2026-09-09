#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Listado de las facturas de Siigo que no se pudieron vincular con una cotizacion,
para que administracion les asigne el numero a mano.

Por que hace falta: el vinculo automatico (server/lib/facturas-vinculo.js) lee el
numero de cotizacion de las observaciones de la factura en Siigo, pero la mayoria
de las facturas ahi traen la ORDEN DE COMPRA del cliente ("O.C. 2179"), que es
otra cosa. Esas quedan sin vincular y hay que resolverlas a mano.

Para no dejar el trabajo entero a la persona, cada factura sale con hasta 3
COTIZACIONES CANDIDATAS, buscadas por dos pistas objetivas:
  - mismo cliente (comparando el nombre normalizado), y
  - monto parecido, teniendo en cuenta que la factura suele traer IVA del 19%
    mientras que el precio de la cotizacion va sin IVA.
Son SUGERENCIAS para revisar, no un vinculo: la columna COTIZACION queda vacia
y la llena una persona.

Uso:
    $env:ADMIN_USERNAME = "admin"
    $env:ADMIN_PASSWORD = "tu-clave"
    python scripts/facturas-sin-cotizacion.py <URL>
"""

import importlib.util
import os
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location(
    'inventario', Path(__file__).resolve().parent / 'listar-elementos-cotizados.py')
inventario = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(inventario)

IVA = 0.19
TOLERANCIA = 0.02  # 2% de diferencia se considera "mismo monto"


def norm(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    s = s.upper()
    # Quita la forma juridica, que se escribe de mil maneras y estorba al comparar
    s = re.sub(r'\b(S\.?A\.?S?\.?|E\.?S\.?P\.?|LTDA\.?|CIA\.?|Y|DE|DEL|LA|EL|LOS|LAS)\b', ' ', s)
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Z0-9 ]', ' ', s)).strip()


def parecido(a, b):
    """Cuanto se parecen dos nombres de cliente, 0 a 1, por palabras compartidas."""
    ta, tb = set(norm(a).split()), set(norm(b).split())
    ta = {t for t in ta if len(t) >= 3}
    tb = {t for t in tb if len(t) >= 3}
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / min(len(ta), len(tb))


def candidatas(factura, cotizaciones):
    """Cotizaciones que podrian corresponder a esta factura."""
    total = float(factura.get('total') or 0)
    # La factura suele venir con IVA; la cotizacion guarda el precio sin IVA.
    base_estimada = total / (1 + IVA) if total else 0
    salida = []
    for c in cotizaciones:
        sim = parecido(factura.get('cliente'), c.get('cliente'))
        if sim < 0.5:
            continue
        precio = float(c.get('precio_venta') or 0)
        if not precio:
            continue
        # Coincide con el total (factura sin IVA) o con la base (factura con IVA)
        d1 = abs(precio - total) / max(precio, total)
        d2 = abs(precio - base_estimada) / max(precio, base_estimada) if base_estimada else 1
        mejor = min(d1, d2)
        if mejor > 0.15:
            continue
        salida.append({
            'numero': c['numero'],
            'precio': precio,
            'dif': mejor,
            'sim_cliente': sim,
            'exacto': mejor <= TOLERANCIA,
            'con_iva': d2 < d1,
        })
    salida.sort(key=lambda x: (x['dif'], -x['sim_cliente']))
    return salida[:3]


COLUMNAS = [
    'COTIZACION', 'NOTAS', 'confianza_sugerencia', 'factura', 'fecha', 'cliente',
    'total', 'saldo', 'estado', 'observaciones_en_siigo',
    'candidata_1', 'candidata_1_precio', 'candidata_2', 'candidata_2_precio',
    'candidata_3', 'candidata_3_precio', 'siigo_invoice_id',
]
ANCHOS = {
    'COTIZACION': 14, 'NOTAS': 22, 'confianza_sugerencia': 12, 'factura': 12,
    'fecha': 11, 'cliente': 34, 'total': 15, 'saldo': 14, 'estado': 10,
    'observaciones_en_siigo': 58, 'candidata_1': 12, 'candidata_1_precio': 14,
    'candidata_2': 12, 'candidata_2_precio': 14, 'candidata_3': 12,
    'candidata_3_precio': 14, 'siigo_invoice_id': 38,
}


def main():
    cfg = inventario.parsear_argumentos(sys.argv[1:])
    usuario, clave = os.environ.get('ADMIN_USERNAME'), os.environ.get('ADMIN_PASSWORD')
    if not usuario or not clave:
        print('Faltan ADMIN_USERNAME / ADMIN_PASSWORD.', file=sys.stderr)
        sys.exit(1)

    api = inventario.ClienteApi(cfg['base_url'])
    api.login(usuario, clave)
    print('Sesion iniciada.')

    facturas = api.get('/api/facturas?desde=2000-01-01&hasta=2030-12-31') or []
    cotizaciones = api.get('/api/cotizaciones') or []
    sin = [f for f in facturas if not f.get('cotizacion_id')]
    print(f'Facturas totales: {len(facturas)} | ya vinculadas: {len(facturas) - len(sin)} | SIN VINCULAR: {len(sin)}')
    print(f'Cotizaciones contra las que se buscan candidatas: {len(cotizaciones)}\n')

    filas = []
    con_exacta = con_alguna = 0
    for f in sorted(sin, key=lambda x: (x.get('fecha') or ''), reverse=True):
        cands = candidatas(f, cotizaciones)
        if cands:
            con_alguna += 1
        exacta = bool(cands and cands[0]['exacto'])
        if exacta:
            con_exacta += 1
        fila = {
            'COTIZACION': None,   # la llena administracion
            'NOTAS': None,
            'confianza_sugerencia': 'MONTO EXACTO' if exacta else ('aproximada' if cands else ''),
            'factura': f.get('numero'),
            'fecha': f.get('fecha'),
            'cliente': f.get('cliente'),
            'total': f.get('total'),
            'saldo': f.get('saldo'),
            'estado': f.get('estado'),
            'observaciones_en_siigo': re.sub(r'\s+', ' ', str(f.get('observaciones') or '')).strip(),
            'siigo_invoice_id': f.get('siigo_invoice_id'),
        }
        for i in range(3):
            fila[f'candidata_{i+1}'] = cands[i]['numero'] if i < len(cands) else None
            fila[f'candidata_{i+1}_precio'] = cands[i]['precio'] if i < len(cands) else None
        filas.append(fila)

    # Primero las que tienen sugerencia de monto exacto: son las mas faciles de
    # confirmar y las que mas rapido reducen la lista.
    orden = {'MONTO EXACTO': 0, 'aproximada': 1, '': 2}
    filas.sort(key=lambda r: (orden[r['confianza_sugerencia']], r['fecha'] or ''), reverse=False)

    carpeta = RAIZ / 'reportes'
    carpeta.mkdir(exist_ok=True)
    marca = datetime.now().strftime('%Y%m%d-%H%M%S')
    ruta_x = carpeta / f'facturas-sin-cotizacion-{marca}.xlsx'
    ruta_c = carpeta / f'facturas-sin-cotizacion-{marca}.csv'

    import csv
    with open(ruta_c, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNAS, delimiter=';', extrasaction='ignore')
        w.writeheader()
        for r in filas:
            w.writerow({k: ('' if r.get(k) is None else r.get(k)) for k in COLUMNAS})

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Asignar cotizacion'
    LLENAR = PatternFill('solid', fgColor='FFF2CC')
    EXACTA = PatternFill('solid', fgColor='C6EFCE')
    APROX = PatternFill('solid', fgColor='FFEB9C')
    ws.append(COLUMNAS)
    for i, c in enumerate(COLUMNAS, 1):
        cel = ws.cell(row=1, column=i)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor='1F4E79')
        cel.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = ANCHOS.get(c, 14)
    ws.freeze_panes = 'C2'
    i_conf = COLUMNAS.index('confianza_sugerencia') + 1
    pesos = {COLUMNAS.index(c) + 1 for c in
             ('total', 'saldo', 'candidata_1_precio', 'candidata_2_precio', 'candidata_3_precio')}
    for r in filas:
        ws.append([r.get(c) for c in COLUMNAS])
        n = ws.max_row
        ws.cell(row=n, column=1).fill = LLENAR
        ws.cell(row=n, column=2).fill = LLENAR
        if r['confianza_sugerencia'] == 'MONTO EXACTO':
            ws.cell(row=n, column=i_conf).fill = EXACTA
        elif r['confianza_sugerencia']:
            ws.cell(row=n, column=i_conf).fill = APROX
        for c in pesos:
            ws.cell(row=n, column=c).number_format = '#,##0'
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta_x)

    print(f'Excel: {ruta_x}')
    print(f'CSV  : {ruta_c}')
    print(f'\nFacturas sin vincular          : {len(filas)}')
    print(f'  con candidata de MONTO EXACTO: {con_exacta}')
    print(f'  con alguna candidata          : {con_alguna}')
    print(f'  sin ninguna pista             : {len(filas) - con_alguna}')


if __name__ == '__main__':
    try:
        main()
    except RuntimeError as e:
        print(f'\nERROR: {e}', file=sys.stderr)
        sys.exit(1)
